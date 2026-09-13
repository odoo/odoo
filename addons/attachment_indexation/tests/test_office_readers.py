import io
import zipfile

from odoo.libs.documents import TEXT, Document, get_known_reader_names, get_readers
from odoo.tests import TransactionCase, tagged

from odoo.addons.attachment_indexation.tools.readers import (
    DOCX,
    MAX_ENTRY_BYTES,
    OPENDOCUMENT,
    PPTX,
    XLSX,
)


def _docx(paragraphs):
    body = "".join(f"<w:p><w:t>{p}</w:t></w:p>" for p in paragraphs)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr(
            "word/document.xml",
            f'<?xml version="1.0"?><w:document xmlns:w="w"><w:body>{body}'
            "</w:body></w:document>",
        )
    return buffer.getvalue()


def _pptx(lines):
    text = "".join(f"<a:t>{line}</a:t>" for line in lines)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr(
            "ppt/slides/slide1.xml",
            f'<?xml version="1.0"?><p:sld xmlns:a="a" xmlns:p="p">{text}</p:sld>',
        )
    return buffer.getvalue()


def _odt(paragraphs):
    body = "".join(f"<text:p>{p}</text:p>" for p in paragraphs)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        zf.writestr(
            "content.xml",
            '<?xml version="1.0"?><office:document-content '
            'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
            f"{body}</office:document-content>",
        )
    return buffer.getvalue()


def _xlsx(rows):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@tagged("post_install", "-at_install")
class TestOfficeReaders(TransactionCase):
    def test_the_four_readers_are_registered(self):
        for name in ("docx_text", "pptx_text", "xlsx_text", "opendocument_text"):
            self.assertIn(name, get_known_reader_names())

    def test_a_word_document_yields_its_text(self):
        document = Document(_docx(["TOTAL 139.86", "CFE"]), name="bill.docx")

        self.assertEqual(document.mimetype, DOCX)
        self.assertIn("139.86", document.text)
        self.assertTrue(document.provides(TEXT))

    def test_a_presentation_yields_its_slide_text(self):
        document = Document(_pptx(["Slide title", "Bullet"]), name="deck.pptx")

        self.assertEqual(document.mimetype, PPTX)
        self.assertIn("Slide title", document.text)

    def test_an_open_document_yields_its_paragraphs(self):
        document = Document(_odt(["ODT PARAGRAPH"]), name="note.odt")

        self.assertIn(document.mimetype, OPENDOCUMENT)
        self.assertIn("ODT PARAGRAPH", document.text)

    def test_a_workbook_yields_its_cells(self):
        document = Document(
            _xlsx([["Ref", "Amount"], ["INV/1", 139.86]]), name="b.xlsx"
        )

        self.assertEqual(document.mimetype, XLSX)
        self.assertIn("139.86", document.text)

    def test_the_mimetype_is_recognised_from_the_bytes_alone(self):
        """`ir.attachment` stores `application/octet-stream` for an upload the
        browser did not label, and the registry keys on the mimetype. A reader
        that could only be reached from a correct declaration would be
        unreachable for exactly the documents the indexer walks every format
        for."""
        document = Document(_docx(["TOTAL 139.86"]), "application/octet-stream", "x")

        self.assertEqual(document.mimetype, DOCX)
        self.assertIn("139.86", document.text)

    def test_the_readers_cost_nothing_so_no_caller_has_to_ask(self):
        for mimetype in (DOCX, PPTX, XLSX, *OPENDOCUMENT):
            with self.subTest(mimetype=mimetype):
                readers = get_readers(mimetype, TEXT)
                self.assertTrue(readers)
                self.assertEqual([r.cost for r in readers], [0] * len(readers))

    def test_an_oversized_entry_is_refused_through_the_registry_too(self):
        """The zip-bomb guard is the reader's, not the indexer's. A document
        reaching the reader through the registry carries no `_INDEX_MAX_BYTES`
        to be bounded by, so the bound has to have a default that holds."""
        padding = b"<!-- " + b"a" * (MAX_ENTRY_BYTES + 1) + b" -->"
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as zf:
            zf.writestr("[Content_Types].xml", "<Types/>")
            zf.writestr(
                "word/document.xml",
                b'<?xml version="1.0"?><w:document xmlns:w="w"><w:body>'
                b"<w:p><w:t>SECRET</w:t></w:p>" + padding + b"</w:body></w:document>",
            )

        self.assertEqual(Document(buffer.getvalue(), DOCX, "bomb.docx").text, "")

    def test_a_caller_may_raise_the_entry_bound(self):
        padding = b"<!-- " + b"a" * (MAX_ENTRY_BYTES + 1) + b" -->"
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as zf:
            zf.writestr("[Content_Types].xml", "<Types/>")
            zf.writestr(
                "word/document.xml",
                b'<?xml version="1.0"?><w:document xmlns:w="w"><w:body>'
                b"<w:p><w:t>SECRET</w:t></w:p>" + padding + b"</w:body></w:document>",
            )
        document = Document(
            buffer.getvalue(),
            DOCX,
            "bomb.docx",
            max_zip_entry_bytes=MAX_ENTRY_BYTES * 2,
        )

        self.assertIn("SECRET", document.text)


@tagged("post_install", "-at_install")
class TestIndexingEdges(TransactionCase):
    def test_an_attachment_with_no_content_is_indexed_rather_than_refused(self):
        """`Document` refuses empty bytes, and an empty attachment is legal.

        Every `_index_*` returned "" for empty content before the indexer read
        through the layer, so nothing here ever built a document out of nothing.
        `_prepare_content_vals` reaches `_get_index_content` for any create, empty raw
        included.
        """
        self.assertFalse(
            self.env["ir.attachment"]._get_index_content(b"", "application/pdf")
        )

        attachment = self.env["ir.attachment"].create(
            {"name": "empty.pdf", "raw": b"", "mimetype": "application/pdf"}
        )

        self.assertFalse(attachment.index_content)

    def test_content_that_is_not_a_document_is_indexed_rather_than_refused(self):
        """A single NUL is non-empty, so it builds a document, and no reader
        claims whatever libmagic makes of it. The interesting part is the NUL
        itself: it must not survive into the column, because `index_content` is
        a text field and PostgreSQL rejects a NUL in one."""
        for content in (b"\x00", b"\x00\x00\x00", b"\xff\xd8\xff\xe0"):
            with self.subTest(content=content):
                indexed = self.env["ir.attachment"]._get_index_content(content, "")
                self.assertNotIn("\x00", indexed or "")

        attachment = self.env["ir.attachment"].create(
            {"name": "nul.bin", "raw": b"\x00"}
        )

        self.assertNotIn("\x00", attachment.index_content or "")
