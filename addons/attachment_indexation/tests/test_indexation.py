# -*- coding: utf-8 -*-

from odoo.tests.common import TransactionCase, tagged
from odoo.tools.misc import file_open
from unittest import skipIf
import io
import openpyxl
import os

directory = os.path.dirname(__file__)

try:
    from pdfminer.pdfinterp import PDFResourceManager
except ImportError:
    PDFResourceManager = None


@tagged('post_install', '-at_install')
class TestCaseIndexation(TransactionCase):

    @skipIf(PDFResourceManager is None, "pdfminer not installed")
    def test_attachment_pdf_indexation(self):
        with file_open(os.path.join(directory, 'files', 'test_content.pdf'), 'rb') as file:
            pdf = file.read()
            text = self.env['ir.attachment']._index(pdf, 'application/pdf')
            self.assertEqual(text, 'TestContent!!', 'the index content should be correct')

    def test_attachment_xlsx_sparse_indexation(self):
        """
        Ensure that extremely sparse Excel files (e.g., spanning 16,384 columns)
        do not cause MemoryErrors by generating billions of empty commas that will
        need to be submitted to .translate() and re.sub() calls.
        """
        # 1. Create a sparse workbook dynamically to avoid storing large/weird files in the repo
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "SparseSheet"

        # Place data at extreme opposite ends of the Excel grid
        ws.cell(row=1, column=1, value="TopLeftData")
        ws.cell(row=150000, column=16384, value="BottomRightData")  # 16,384 is the max Excel column limit

        # Populate the first column across many rows.
        # Because each row has at least one value, `any(row)` code evaluates to True.
        # If not handled properly, the indexing will then attempt to process
        # the remaining 16,383 empty cells per row.
        # From local tests, 50K populated rows should be enough to trigger memory errors
        # in a default odoo-bin run.
        for r in range(2, 50000):
            ws.cell(row=r, column=1, value=f"Data {r}")

        # Save to an in-memory byte stream
        excel_stream = io.BytesIO()
        wb.save(excel_stream)
        excel_binary_data = excel_stream.getvalue()

        # Trigger the standard indexation routing
        mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        indexed_text = self.env['ir.attachment']._index(excel_binary_data, mimetype)

        # If the indexing triggers a memory error, it should fail before this and
        # return None for `indexed_text`` variable that will fail the assertions
        self.assertIsNotNone(indexed_text, 'No indexed data returned, probably failed while indexing')
        self.assertIn('TopLeftData', indexed_text, 'The index should contain the first cell data')
        self.assertIn('BottomRightData', indexed_text, 'The index should contain the far-right cell data')
        self.assertIn('SparseSheet', indexed_text, 'The index should contain the sheet name')
