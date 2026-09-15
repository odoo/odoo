import io
import unittest

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from odoo import Command
from odoo.tests import HttpCase, tagged


@unittest.skipIf(load_workbook is None, "openpyxl not available")
class TestAssetTemplate(HttpCase):
    def test_download_asset_template(self):
        self.authenticate("admin", "admin")
        company = self.env.company
        account_model = self.env["account.account"].with_company(company)

        account_model.search(
            [
                (
                    "account_type",
                    "in",
                    [
                        "asset_fixed",
                        "asset_non_current",
                        "expense_depreciation",
                        "expense",
                    ],
                )
            ]
        ).write({"active": False})

        account_model.create(
            {
                "name": "Fixed Asset Account",
                "code": "FA001",
                "account_type": "asset_non_current",
            }
        )
        account_model.create(
            {
                "name": "Expense Account",
                "code": "EX001",
                "account_type": "expense",
            }
        )

        self.env["account.journal"].with_company(company).create(
            {
                "name": "Miscellaneous Operations",
                "code": "MISC-test",
                "type": "general",
            }
        )

        response = self.url_open(f"/web/binary/download_asset_template/{company.id}")
        self.assertEqual(response.status_code, 200)

        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook["Assets"]

        expected_rows = [
            (
                "Computer 1",
                "800.00",
                "01-01-2025",
                "",
                "Straight Line",
                "4",
                "Years",
                "",
                "No Prorata",
                "01-01-2025",
                "200.00",
                "0.00",
                company.name,
                "FA001 Fixed Asset Account",
                "FA001 Fixed Asset Account",
                "EX001 Expense Account",
                "Miscellaneous Operations",
            ),
            (
                "Computer 2",
                "25000.00",
                "03-15-2025",
                "",
                "Declining",
                "5",
                "Years",
                "2",
                "Based on days per period",
                "03-15-2025",
                "0.00",
                "2000.00",
                company.name,
                "FA001 Fixed Asset Account",
                "FA001 Fixed Asset Account",
                "EX001 Expense Account",
                "Miscellaneous Operations",
            ),
            (
                "Machine A",
                "15000.00",
                "09-01-2024",
                "",
                "Straight Line",
                "10",
                "Years",
                "",
                "Constant Periods",
                "09-01-2024",
                "1500.00",
                "500.00",
                company.name,
                "FA001 Fixed Asset Account",
                "FA001 Fixed Asset Account",
                "EX001 Expense Account",
                "Miscellaneous Operations",
            ),
            (
                "Machine B",
                "100000.00",
                "06-20-2025",
                "",
                "Straight Line",
                "60",
                "Months",
                "",
                "No Prorata",
                "06-20-2025",
                "10000.00",
                "10000.00",
                company.name,
                "FA001 Fixed Asset Account",
                "FA001 Fixed Asset Account",
                "EX001 Expense Account",
                "Miscellaneous Operations",
            ),
        ]

        actual_rows = [
            tuple("" if v is None else str(v) for v in actual_row[:17])
            for actual_row in sheet.iter_rows(min_row=2, values_only=True)
        ]
        self.assertEqual(actual_rows, expected_rows)


@tagged("post_install", "-at_install")
class TestAssetTemplateAccess(HttpCase):
    def test_a_company_the_user_cannot_see_is_refused(self):
        other = self.env["res.company"].create({"name": "Secret Holdings SA"})
        admin = self.env.ref("base.user_admin")
        admin.write(
            {
                "company_ids": [Command.set([self.env.company.id])],
                "company_id": self.env.company.id,
            }
        )
        self.env.flush_all()
        self.authenticate("admin", "admin")

        response = self.url_open(f"/web/binary/download_asset_template/{other.id}")

        self.assertEqual(response.status_code, 403)
        self.assertNotEqual(response.content[:2], b"PK", "no workbook may be served")
        self.assertNotIn("Secret Holdings", response.text)

    def test_a_portal_user_gets_no_workbook(self):
        self.env["res.users"].create(
            {
                "name": "Portal Probe",
                "login": "portal_probe",
                "password": "portal_probe",
                "group_ids": [Command.set([self.env.ref("base.group_portal").id])],
            }
        )
        self.authenticate("portal_probe", "portal_probe")

        response = self.url_open(
            f"/web/binary/download_asset_template/{self.env.company.id}"
        )

        self.assertEqual(response.status_code, 403)
        self.assertNotEqual(response.content[:2], b"PK")
